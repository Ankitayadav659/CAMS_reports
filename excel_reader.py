"""
Streaming Excel reader.

* openpyxl read-only  -> true streaming (lowest RAM)
* python-calamine     -> very fast; used automatically for moderate files

Both expose the same tiny interface:   header()  /  rows()  /  close()
"""
import os
import zipfile

os.environ.setdefault('RUST_BACKTRACE', '0')   # keep calamine panics quiet (we fall back to openpyxl)
from datetime import date, datetime, time, timedelta

import config
from reporting import UserError, log


class SheetReader:
    def __init__(self, path, sheet_name=None, engine=None):
        if not os.path.isfile(path):
            raise UserError(f'File not found: {path}')
        if os.path.splitext(path)[1].lower() not in ('.xlsx', '.xlsm'):
            raise UserError(f'Only .xlsx / .xlsm files are supported: {os.path.basename(path)}')
        engine = engine or config.READER
        if engine == 'auto':
            mb = os.path.getsize(path) / 1_048_576
            engine = 'calamine' if mb <= config.CALAMINE_MAX_FILE_MB else 'openpyxl'
        if engine == 'calamine':
            try:
                import python_calamine  # noqa: F401
            except ImportError:
                engine = 'openpyxl'
        self.engine = engine
        self.path = path
        self._wb = None
        self._it = None
        self.total_rows = None
        try:
            try:
                self._open(sheet_name)
            except (KeyboardInterrupt, SystemExit, UserError):
                raise
            except BaseException as e:               # calamine (Rust) panics are BaseException subclasses
                if self.engine != 'calamine':
                    raise
                log(f'         calamine could not read this file ({type(e).__name__}); using openpyxl instead')
                self.engine = 'openpyxl'
                self._open(sheet_name)
        except UserError:
            raise
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
            raise UserError(f'"{os.path.basename(path)}" is not a valid / readable Excel file ({e}).')
        except Exception as e:                       # openpyxl InvalidFileException, calamine errors
            raise UserError(f'Could not open "{os.path.basename(path)}": {e}')

    # ------------------------------------------------------------------
    def _open(self, sheet_name):
        if self.engine == 'calamine':
            from python_calamine import CalamineWorkbook
            self._wb = CalamineWorkbook.from_path(self.path)
            names = self._wb.sheet_names
            if sheet_name is not None and sheet_name not in names:
                raise UserError(f'Worksheet "{sheet_name}" not found. Sheets present: {names}')
            sh = self._wb.get_sheet_by_name(sheet_name) if sheet_name else self._wb.get_sheet_by_index(0)
            self.total_rows = sh.height
            self._it = iter(sh.iter_rows())
        else:
            from openpyxl import load_workbook
            self._wb = load_workbook(self.path, read_only=True, data_only=True)
            names = self._wb.sheetnames
            if sheet_name is not None and sheet_name not in names:
                self._wb.close()
                raise UserError(f'Worksheet "{sheet_name}" not found. Sheets present: {names}')
            ws = self._wb[sheet_name] if sheet_name else self._wb.worksheets[0]
            try:
                ws.reset_dimensions()               # ignore a wrong <dimension> tag
            except Exception:
                pass
            self.total_rows = ws.max_row if ws.max_row and ws.max_row > 1 else None
            self._it = ws.iter_rows(values_only=True)

    def header(self):
        try:
            row = next(self._it)
        except StopIteration:
            raise UserError(f'"{os.path.basename(self.path)}" has no data (empty worksheet).')
        return ['' if v is None else str(v).strip() for v in row]

    def rows(self):
        return self._it

    def close(self):
        try:
            if self._wb is not None and hasattr(self._wb, 'close'):
                self._wb.close()
        except Exception:
            pass


def open_sheet(path, sheet_name=None, engine=None):
    r = SheetReader(path, sheet_name, engine)
    log(f'         Reader: {r.engine}' + (' (fast)' if r.engine == 'calamine' else ' (streaming, low RAM)'))
    return r


# ───────────────────────── value helpers ─────────────────────────
_NBSP = '\u00a0'


def cell_to_str(v):
    """Mimic pandas dtype=str: None/'' -> None, numbers without trailing .0 ."""
    if v is None:
        return None
    if isinstance(v, str):
        return v if v != '' else None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        if v != v:
            return None
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, datetime):
        return str(v)
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


def norm_id(v):
    """Normalise an identifier for EXACT matching.
    Keeps leading zeros; strips outer/NBSP spaces; 1015586806.0 -> '1015586806'."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        if v != v:
            return ''
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, (datetime, date, time, timedelta)):
        return str(v).strip()
    return str(v).replace(_NBSP, ' ').strip()


def norm_name(v):
    """Investor-name normalisation for validation: case-insensitive, whitespace-collapsed."""
    return ' '.join(norm_id(v).split()).casefold()


def norm_amc(v):
    return norm_id(v).upper()
