"""
UPDATE EXISTING INVESTOR SHEETS

Take investor Excel files that already exist (for example the "final investors" sheets, with or without the
Brokerage column) and write NEW copies that have
    * the Brokerage column recalculated (Previous BRKAGE_RAT − Current BRKAGE_RAT, TRADE_DATE order),
    * the Rate Change column,
    * the simple "Rate Summary" sheet in front of every investor sheet.

The originals are never changed.  It uses exactly the same writer as Step 2, so the result looks identical.
A sheet is treated as an investor sheet when it has the columns TRADE_DATE and BRKAGE_RAT; every other sheet
(an old Rate Summary, an index, ...) is dropped and the Rate Summary is rebuilt.  Re-running on the output is safe.
"""
import os
import time

import xlsxwriter
from openpyxl import load_workbook

import config as C
from database import TempStore
from excel_reader import SheetReader, norm_id
from excel_writer import Formats, SheetNamer, abort_workbook, write_group_sheets
from investor_engine import DateParser, Group, parse_rate
from reporting import UserError, log, progress

DROP_COLS = {C.BROKERAGE_COL, C.RATE_CHANGE_COL}          # recalculated, so not copied from the old file


def _sheet_names(path):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise UserError(f'not a valid / readable Excel file ({type(e).__name__})')
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _first(row, i):
    return norm_id(row[i]) if i is not None and i < len(row) else ''


def process_file(path, out_path):
    """Returns (data_sheets_written, skipped_sheet_names)."""
    names = _sheet_names(path)
    parse_date = DateParser()
    store = TempStore(os.path.dirname(os.path.abspath(out_path)))
    groups, headers, skipped = [], {}, []
    try:
        for name in names:
            try:
                rd = SheetReader(path, name, 'openpyxl')
            except UserError as e:
                raise UserError(f'sheet "{name}": {e}')
            try:
                hdr = rd.header()
                if C.DATE_COL not in hdr or C.BROK_COL not in hdr:
                    skipped.append(name)
                    continue
                keep = [i for i, h in enumerate(hdr) if h not in DROP_COLS]
                src_hdr = [hdr[i] for i in keep]
                idx = {h: k for k, h in reversed(list(enumerate(src_hdr)))}          # first occurrence wins
                d_i, b_i = idx[C.DATE_COL], idx[C.BROK_COL]
                meta = {k: idx.get(k) for k in (C.INV_COL, C.AMC_COL, C.SCHEME_COL, C.FOLIO_COL, C.TXN_TYPE_COL)}
                found = {k: '' for k in meta}
                gid = len(groups)
                for seq, row in enumerate(rd.rows(), start=1):
                    if all(v is None for v in row):
                        continue
                    r = [row[i] if i < len(row) else None for i in keep]
                    for k, i in meta.items():
                        if not found[k]:
                            found[k] = _first(r, i)
                    rate, rstat = parse_rate(r[b_i])
                    store.add(gid, seq, parse_date(r[d_i]), rate, rstat, r)
                g = Group(gid, found[C.INV_COL] or name, found[C.AMC_COL], found[C.SCHEME_COL],
                          found[C.FOLIO_COL], found[C.TXN_TYPE_COL])
                g.file = name                                                           # original sheet name
                groups.append(g)
                headers[gid] = (src_hdr, d_i, b_i)
            finally:
                rd.close()
        if not groups:
            return 0, skipped
        store.finalize()
        counts = store.group_counts()
        namer = SheetNamer()
        live = []
        for g in groups:
            g.count = counts.get(g.gid, (0,))[0]
            if not g.count:
                skipped.append(g.file)
                continue
            g.n_sheets = -(-g.count // C.STEP2_ROWS_PER_SHEET)
            first = namer.make(g.file)                                                  # keep the old sheet name
            g.sheets = [first] + [namer.make(first, suffix=f'p{i}') for i in range(2, g.n_sheets + 1)]
            live.append(g)
        if not live:
            return 0, skipped
        for g in live:                                                                  # summary names after data names
            g.summary = namer.make('Rate Summary' if len(live) == 1 else 'Summary ' + g.file)

        wb = xlsxwriter.Workbook(out_path, {'constant_memory': True, 'strings_to_numbers': False,
                                            'use_zip64': True, 'tmpdir': os.path.dirname(os.path.abspath(out_path))})
        try:
            fm = Formats(wb)
            for g in live:
                hdr, d_i, b_i = headers[g.gid]
                write_group_sheets(wb, fm, g, store, hdr, d_i, b_i)
            wb.close()
        except BaseException:
            abort_workbook(wb)
            try:
                os.remove(out_path)
            except OSError:
                pass
            raise
        return len(live), skipped
    finally:
        store.close()


def _collect(in_path):
    if os.path.isfile(in_path):
        return [in_path]
    if not os.path.isdir(in_path):
        raise UserError(f'Folder / file not found: {in_path}')
    files = [os.path.join(in_path, f) for f in sorted(os.listdir(in_path), key=str.casefold)
             if f.lower().endswith(('.xlsx', '.xlsm')) and not f.startswith('~$')]
    if not files:
        raise UserError(f'No Excel (.xlsx) files found in "{in_path}".')
    return files


def run_update(in_path, out_folder):
    t0 = time.time()
    if not out_folder:
        raise UserError('Please choose an output folder.')
    in_dir = in_path if os.path.isdir(in_path) else os.path.dirname(os.path.abspath(in_path))
    if os.path.abspath(in_dir) == os.path.abspath(out_folder):
        raise UserError('The output folder must be different from the folder with your original files, '
                        'so the originals are never overwritten.')
    files = _collect(in_path)
    try:
        os.makedirs(out_folder, exist_ok=True)
    except OSError as e:
        raise UserError(f'Cannot create the output folder "{out_folder}": {e}')
    log(f'Found {len(files)} Excel file(s). Updated copies will be saved in:\n   {out_folder}\n')

    ok = sheets = 0
    skipped_files, failed = [], []
    for n, f in enumerate(files, start=1):
        base = os.path.basename(f)
        progress(100 * (n - 1) / len(files), f'Updating {n} / {len(files)}:  {base}')
        out = os.path.join(out_folder, base)
        try:
            done, skipped = process_file(f, out)
        except PermissionError:
            failed.append((base, 'file is open in Excel or not writable — close it and run again'))
            log(f'  ✗ {base}  —  {failed[-1][1]}')
            continue
        except UserError as e:
            failed.append((base, str(e)))
            log(f'  ✗ {base}  —  {e}')
            continue
        except Exception as e:
            failed.append((base, f'{type(e).__name__}: {e}'))
            log(f'  ✗ {base}  —  {type(e).__name__}: {e}')
            continue
        if done:
            ok += 1
            sheets += done
            log(f'  ✓ {base}  —  {done} investor sheet(s) updated')
        else:
            skipped_files.append(base)
            log(f'  ⚠ {base}  —  skipped (no sheet with TRADE_DATE and BRKAGE_RAT columns)')
    progress(100, 'Done')
    if not ok:
        raise UserError('No file could be updated. ' + (f'First problem: {failed[0][0]}: {failed[0][1]}' if failed else
                        'None of the files has a sheet with TRADE_DATE and BRKAGE_RAT columns.'))
    log(f'\nFinished in {time.time() - t0:,.1f} s  —  updated {ok} file(s), {len(skipped_files)} skipped, {len(failed)} failed.')
    return {'files': ok, 'sheets': sheets, 'skipped': len(skipped_files), 'failed': len(failed),
            'failed_list': failed, 'output': out_folder}
